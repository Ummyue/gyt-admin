#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""DOCX 成品校验：结构 / 图片 / 品牌字样 / 字体 / 域"""
import sys, os, zipfile, re, json
from docx import Document
from docx.shared import Pt

DOCX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "数字供应链平台-产品需求文档汇编.docx")

print("=" * 60)
print("文件:", DOCX)
print("大小: %.2f MB" % (os.path.getsize(DOCX) / 1024 / 1024))

# 1) zip 完整性
z = zipfile.ZipFile(DOCX)
bad = z.testzip()
print("ZIP 完整性:", "OK" if bad is None else f"损坏 {bad}")

# 2) 品牌字样（全文，含所有 part）
BAD = ["豫港通", "中豫港通", "港通供应链", "中豫", "豫港"]
leaks = {}
for n in z.namelist():
    if n.endswith((".xml", ".rels")):
        try:
            t = z.read(n).decode("utf-8", "ignore")
        except Exception:
            continue
        hit = [k for k in BAD if k in t]
        if hit:
            leaks[n] = hit
print("品牌字样扫描:", "无泄漏 ✅" if not leaks else f"❌ {leaks}")

# 3) 结构读回
d = Document(DOCX)
heads = [(p.style.name, p.text.strip()) for p in d.paragraphs
         if p.style.name.startswith("Heading") and p.text.strip()]
h1 = [t for s, t in heads if s == "Heading 1"]
h2 = [t for s, t in heads if s == "Heading 2"]
h3 = [t for s, t in heads if s == "Heading 3"]
print(f"段落数: {len(d.paragraphs)} | 表格数: {len(d.tables)} | 图片数: {len(d.inline_shapes)}")
print(f"Heading1: {len(h1)} | Heading2: {len(h2)} | Heading3: {len(h3)}")
print("H1 列表:", " / ".join(h1))
print("模块章数(4.x):", len([t for t in h2 if t.startswith("4.")]))
print("小节(4.x.y):", len([t for t in h3 if re.match(r"^4\.\d+\.\d+", t)]))

# 4) 字体检查
def eastasia(style):
    try:
        return style.element.rPr.rFonts.get(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia")
    except Exception:
        return None
print("Normal eastAsia 字体:", eastasia(d.styles["Normal"]))
for i in (1, 2, 3):
    print(f"Heading {i} eastAsia 字体:", eastasia(d.styles[f"Heading {i}"]))

# 5) 域（目录/页码）
doc_xml = z.read("word/document.xml").decode("utf-8", "ignore")
print("TOC 域:", "有 ✅" if "TOC \\o" in doc_xml or 'TOC \\o' in doc_xml else "无 ❌")
foot_xml = "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.startswith("word/footer"))
print("页脚 PAGE 域:", "有 ✅" if "PAGE" in foot_xml else "无 ❌")
print("NUMPAGES 域:", "有 ✅" if "NUMPAGES" in foot_xml else "无 ❌")

# 6) 分节与页眉
print("节数:", len(d.sections))
for i, s in enumerate(d.sections):
    hp = s.header.paragraphs[0].text.strip() if s.header.paragraphs else ""
    print(f"  节{i+1}: 页眉='{hp[:44]}' 页宽={s.page_width.cm:.1f}cm 左边距={s.left_margin.cm:.1f}cm")

# 7) 图片清单
media = [n for n in z.namelist() if n.startswith("word/media/")]
print("内嵌媒体文件:", len(media))

# 8) 空表检测
empty = [i for i, t in enumerate(d.tables) if len(t.rows) < 2]
print("空表(仅表头):", empty if empty else "无 ✅")

# 9) 章节完整性
need = ["1　文档说明", "2　产品概述", "3　全局规范与共性约定",
        "4　各模块需求说明", "附录 A　页面索引", "附录 B　术语与缩写"]
FORBID = ["5　跨模块待确认问题清单"]
miss = [n for n in need if n not in h1]
print("必备章节:", "齐全 ✅" if not miss else f"缺 {miss}")
present_forbidden = [n for n in FORBID if n in h1]
print("已删除章节残留:", "无 ✅" if not present_forbidden else f"❌ 仍存在 {present_forbidden}")
# 正文里不得再出现指向第 5 章的引用
raw = z.read("word/document.xml").decode("utf-8", "ignore")
body_txt = re.sub(r"<[^>]+>", "", raw)
stale = [k for k in ["第 5 章", "跨模块待确认问题清单"] if k in body_txt]
print("失效交叉引用:", "无 ✅" if not stale else f"❌ {stale}")
